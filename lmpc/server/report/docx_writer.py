"""The same sections, as an editable document.

python-docx has no true watermark, so a FIELD copy gets a bold banner paragraph and a
page header instead. That is an approximation, and it is described as one rather than
claimed to be equivalent.
"""
from __future__ import annotations

import io


def render(sections: list[dict], *, kind: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    if kind == "FIELD":
        for section in document.sections:
            section.header.paragraphs[0].text = "FIELD COPY — NOT LEGALLY FINALISED"
    for item in sections:
        _RENDERERS[item["kind"]](document, item, Pt)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _cover(document, section, Pt) -> None:
    document.add_heading(section["title"], 0)
    document.add_paragraph("Legal Metrology (Packaged Commodities) Rules, 2011")
    if "not legally" in section["kind_label"].lower():
        document.add_paragraph(section["kind_label"].upper()).runs[0].bold = True
    verdict = document.add_paragraph(section["verdict_word"])
    verdict.runs[0].bold = True
    verdict.runs[0].font.size = Pt(15)
    document.add_paragraph(str(section["verdict"] or ""))
    for label, value in section["rows"]:
        document.add_paragraph(f"{label}: {value}")


def _counts(document, section, _Pt) -> None:
    document.add_heading(section["title"], 1)
    for outcome, count in sorted(section["counts"].items()):
        document.add_paragraph(
            f"{count} — {section['words'].get(outcome, outcome)}", style="List Bullet")
    if section["note"]:
        document.add_paragraph(section["note"])


def _table(document, section, _Pt) -> None:
    if not section["rows"]:
        return
    document.add_heading(section["title"], 1)
    table = document.add_table(rows=1, cols=len(section["columns"]))
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, section["columns"]):
        cell.text = title
        cell.paragraphs[0].runs[0].bold = True
    for row in section["rows"]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)


def _images(document, section, _Pt) -> None:
    from docx.shared import Inches

    document.add_heading(section["title"], 1)
    for item in section["images"]:
        if item.get("raw"):
            try:
                document.add_picture(io.BytesIO(item["raw"]), width=Inches(2.0))
            except Exception:                               # noqa: BLE001
                pass
        document.add_paragraph(f"{item['panel']} — SHA-256 {item['sha256']}")
    document.add_paragraph(section["note"])


def _notapplicable(document, section, _Pt) -> None:
    if not section["items"]:
        return
    document.add_heading(section["title"], 1)
    document.add_paragraph(
        "These were evaluated and found not to apply to this package: "
        + ", ".join(str(item) for item in section["items"]) + ".")


def _list(document, section, _Pt) -> None:
    document.add_heading(section["title"], 1)
    for item in section["items"]:
        document.add_paragraph(str(item), style="List Bullet")


def _notes(document, section, _Pt) -> None:
    document.add_heading(section["title"], 1)
    for note in section["notes"]:
        document.add_paragraph(note["body"])
        document.add_paragraph(
            f"{note.get('author') or ''} · {note.get('created_at') or ''}")


def _signoff(document, section, _Pt) -> None:
    document.add_heading(section["title"], 1)
    for label in ("Name", "Signature", "Date"):
        document.add_paragraph(f"{label}: ______________________________")


def _integrity(document, section, _Pt) -> None:
    document.add_heading(section["title"], 1)
    for label, value in section["rows"]:
        if value:
            document.add_paragraph(f"{label}: {value}")


_RENDERERS = {"cover": _cover, "counts": _counts, "table": _table, "images": _images,
              "notapplicable": _notapplicable, "list": _list, "notes": _notes, "signoff": _signoff,
              "integrity": _integrity}
