"""Versioned PDF and DOCX reports rendered from one canonical verdict snapshot."""
from __future__ import annotations

import hashlib
import html
import io
import json
from datetime import UTC, datetime
from typing import Any

from lmpc.engine import ocr_models
from lmpc.engine.extract import ACCUSE, FLOOR, LEGIBLE, MARGIN
from lmpc.engine.ocr import CAP_RATIO

from ..api.errors import ApiError
from .object_store import ObjectStore


class ReportService:
    def __init__(self, scans, objects: ObjectStore, rulepack: dict, *, max_edge: int,
                 git_sha: str = "", container_digest: str = ""):
        self.scans = scans
        self.objects = objects
        self.rulepack = rulepack
        self.max_edge = max_edge
        self.git_sha = git_sha
        self.container_digest = container_digest

    def finalize(self, scan_id: str) -> dict:
        scan = self.scans.get(scan_id)
        if not scan.latest_evaluations():
            raise ApiError("E_CONFLICT", "the scan has no completed evaluation batch")
        if any(item["outcome"] == "SYSTEM_ERROR" for item in scan.latest_evaluations()):
            raise ApiError("E_CONFLICT", "unresolved SYSTEM_ERROR results block finalization")
        version = self.scans.next_report_version(scan_id)
        snapshot = _snapshot(scan, self.rulepack)
        canonical = json.dumps(snapshot, sort_keys=True, separators=(",", ":"),
                               ensure_ascii=False).encode()
        content_sha256 = hashlib.sha256(canonical).hexdigest()
        manifest = self._manifest(scan, content_sha256)
        pdf = _pdf(snapshot, version, content_sha256)
        docx = _docx(snapshot, version, content_sha256)
        prefix = f"reports/{scan_id}/{version}"
        pdf_key, docx_key = f"{prefix}/report.pdf", f"{prefix}/report.docx"
        self.objects.put_immutable(
            pdf_key, pdf, "application/pdf", hashlib.sha256(pdf).hexdigest())
        self.objects.put_immutable(
            docx_key, docx,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            hashlib.sha256(docx).hexdigest())
        return self.scans.save_report(
            scan_id, version=version, overall_status=scan.overall,
            pdf_storage_key=pdf_key, docx_storage_key=docx_key,
            content_sha256=content_sha256, manifest=manifest)

    def download(self, report_id: str, format_: str) -> tuple[bytes, str, str]:
        report = self.scans.get_report(report_id)
        if format_ == "pdf":
            return self.objects.read(report["pdf_storage_key"]), "application/pdf", "report.pdf"
        if format_ == "docx":
            media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            return self.objects.read(report["docx_storage_key"]), media, "report.docx"
        raise ApiError("E_VALIDATION", "format must be pdf or docx")

    def _manifest(self, scan, content_sha256: str) -> dict:
        models = [{"name": name, "sha256": digest} for name, digest in
                  sorted(ocr_models.entry().get("devanagari", {}).items())]
        return {
            "scan_id": scan.id, "evaluation_batch": scan.batch,
            "evaluated_at": datetime.now(UTC).isoformat(),
            "content_sha256": content_sha256,
            "rulepack": {"version": scan.rulepack_version,
                         "sha256": scan.rulepack_sha256},
            "models": models,
            "params": {"MAX_EDGE": self.max_edge, "CAP_RATIO": CAP_RATIO,
                       "FLOOR": FLOOR, "MARGIN": MARGIN,
                       "LEGIBLE": LEGIBLE, "ACCUSE": ACCUSE},
            "code": {"git_sha": self.git_sha, "container_digest": self.container_digest},
            "inputs": [{"storage_key": image.storage_key, "sha256": image.sha256,
                        "max_edge_used": image.max_edge_used} for image in scan.images],
        }


def _snapshot(scan, pack: dict) -> dict[str, Any]:
    return {
        "identification": {"scan_id": scan.id, "category": scan.category,
                           "mode": scan.mode},
        "overall": scan.overall,
        "captured_at": scan.captured_at,
        "coverage": {"asserted": scan.coverage_asserted, "panels": scan.panels},
        "declarations": scan.latest_declarations(),
        "evaluations": scan.latest_evaluations(),
        "disclosures": [item["text"] for item in
                        pack["currency"]["acknowledged_gaps"]],
        "rulepack": {"version": scan.rulepack_version, "sha256": scan.rulepack_sha256,
                     "current_to": pack["currency"]["newest_instrument"]},
        "images": [{"panel": image.panel_label, "sha256": image.sha256,
                    "storage_key": image.storage_key} for image in scan.images],
    }


def _pdf(snapshot: dict, version: int, content_hash: str) -> bytes:
    from weasyprint import HTML

    return HTML(string=_html(snapshot, version, content_hash)).write_pdf()


def _html(snapshot: dict, version: int, content_hash: str) -> str:
    esc = lambda value: html.escape(str(value))
    rows = "".join(
        "<tr>" + "".join(f"<td>{esc(value)}</td>" for value in (
            item["check"], item["clause"], item["outcome"], item["reason"],
            json.dumps(item.get("evidence") or {}, ensure_ascii=False),
            _citation(item.get("citation") or {}))) + "</tr>"
        for item in snapshot["evaluations"])
    disclosures = "".join(f"<li>{esc(item)}</li>" for item in snapshot["disclosures"])
    evidence = "".join(
        f"<li>{esc(item['panel'])}: <span class='hash'>{esc(item['sha256'])}</span></li>"
        for item in snapshot["images"])
    counts = {}
    for item in snapshot["evaluations"]:
        counts[item["outcome"]] = counts.get(item["outcome"], 0) + 1
    coverage = snapshot["coverage"]
    return f"""<!doctype html><meta charset="utf-8"><style>
body{{font:9pt sans-serif;color:#17212b}} h1{{font-size:18pt}} table{{border-collapse:collapse}}
th,td{{border:1px solid #9aa5b1;padding:5px;vertical-align:top}} th{{background:#eef2f6}}
.hash{{font:8pt monospace;overflow-wrap:anywhere}} @page{{size:A4;margin:14mm;
@bottom-center{{content:"Page " counter(page) " of " counter(pages)}}}}
</style><h1>LMPC Compliance Report — {esc(snapshot['overall'])}</h1>
<h2>Identification</h2><p>Scan {esc(snapshot['identification']['scan_id'])} ·
category {esc(snapshot['identification']['category'])} ·
mode {esc(snapshot['identification']['mode'])} · report version {version}</p>
<h2>Overall status</h2><p>{esc(snapshot['overall'])} · {esc(counts)}</p>
<h2>Scan metadata</h2><p>Captured {esc(snapshot['captured_at'])}</p>
<h2>Coverage</h2><p>Asserted: {esc(coverage['asserted'])}; panels:
{esc(', '.join(coverage['panels']))}</p>
<h2>Findings</h2><table><thead><tr><th>Check</th><th>Clause</th><th>Verdict</th><th>Reason</th>
<th>Measured / required</th><th>Authority</th></tr></thead><tbody>{rows}</tbody></table>
<h2>Evidence</h2><ul>{evidence}</ul><h2>Officer notes</h2><p>—</p>
<h2>Reviewer sign-off</h2><p>Name: ____________________ Date: __________</p>
<h2>Chain-gap disclosures</h2><ul>{disclosures}</ul><h2>Integrity</h2>
<p>Rulepack {esc(snapshot['rulepack']['version'])}</p>
<p class="hash">Rulepack SHA-256: {esc(snapshot['rulepack']['sha256'])}<br>
Report content SHA-256: {content_hash}</p>"""


def _docx(snapshot: dict, version: int, content_hash: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(f"LMPC Compliance Report — {snapshot['overall']}", 0)
    document.add_paragraph(
        f"Scan {snapshot['identification']['scan_id']} · report version {version} · "
        f"captured {snapshot['captured_at']}")
    document.add_heading("Overall status", 1)
    document.add_paragraph(str(snapshot["overall"]))
    document.add_heading("Coverage", 1)
    coverage = snapshot["coverage"]
    document.add_paragraph(
        f"Asserted: {coverage['asserted']}; panels: {', '.join(coverage['panels'])}")
    document.add_heading("Findings", 1)
    table = document.add_table(rows=1, cols=6)
    for cell, title in zip(table.rows[0].cells,
                           ("Check", "Clause", "Verdict", "Reason",
                            "Measured / required", "Authority")):
        cell.text = title
    for item in snapshot["evaluations"]:
        cells = table.add_row().cells
        values = (item["check"], item["clause"], item["outcome"], item["reason"],
                  json.dumps(item.get("evidence") or {}, ensure_ascii=False),
                  _citation(item.get("citation") or {}))
        for cell, value in zip(cells, values):
            cell.text = str(value)
    document.add_heading("Evidence", 1)
    for item in snapshot["images"]:
        document.add_paragraph(f"{item['panel']}: SHA-256 {item['sha256']}")
    document.add_heading("Officer notes", 1)
    document.add_paragraph("—")
    document.add_heading("Reviewer sign-off", 1)
    document.add_paragraph("Name: ____________________ Date: __________")
    document.add_heading("Chain-gap disclosures", 1)
    for disclosure in snapshot["disclosures"]:
        document.add_paragraph(disclosure, style="List Bullet")
    document.add_heading("Integrity", 1)
    document.add_paragraph(f"Rulepack {snapshot['rulepack']['version']}")
    document.add_paragraph(f"Rulepack SHA-256: {snapshot['rulepack']['sha256']}")
    document.add_paragraph(f"Report content SHA-256: {content_hash}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _citation(value: dict) -> str:
    return " · ".join(str(value[key]) for key in ("gsr", "dated", "page") if value.get(key))
