"""PDF/DOCX finalization from one versioned verdict snapshot."""
from __future__ import annotations

import hashlib
import io

import pytest
from docx import Document

from lmpc.engine.model import Token
from lmpc.lawc.build import load
from lmpc.server.api.errors import ApiError
from lmpc.server.svc.evidence import EvidenceImage
from lmpc.server.svc.object_store import LocalObjectStore
from lmpc.server.svc.pipeline import Pipeline
from lmpc.server.svc.reporting import ReportService
from lmpc.server.svc.scan_service import MemoryStore


def _evaluated(tmp_path):
    raw = b"report-source-image"
    digest = hashlib.sha256(raw).hexdigest()
    image = EvidenceImage(
        filename="front.jpg", media_type="image/jpeg", data=b"", sha256=digest,
        width_px=800, height_px=600, storage_key=f"images/{digest}.jpg",
        panel_label="FRONT")
    objects = LocalObjectStore(tmp_path)
    objects.put_immutable(image.storage_key, raw, image.media_type, digest)
    scans = MemoryStore()
    scan, _ = scans.create(
        client_uuid="20000000-0000-4000-8000-000000000001",
        captured_at="2026-09-07", mode="PHYSICAL_PACKAGE", category="FOOD",
        coverage_asserted=False, panels=["FRONT"], images=[image])

    def reader(_data, *, panel, **_kwargs):
        return [Token("MRP Rs. 45.00 (incl. of all taxes)", 10, 20, 300, 20,
                      conf=0.99, panel=panel)]

    pack = load()
    Pipeline(scans, objects, pack, 1800, reader=reader).process(scan.id)
    return scans, objects, pack, scan.id


def _document_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def test_both_formats_are_stored_with_the_same_finding_set_and_content_hash(tmp_path):
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    reports = ReportService(scans, objects, pack, max_edge=1800, git_sha="abc123")
    report = reports.finalize(scan_id)
    pdf, _, _ = reports.download(report["id"], "pdf")
    docx, _, _ = reports.download(report["id"], "docx")

    assert pdf.startswith(b"%PDF-") and docx.startswith(b"PK")
    text = _document_text(docx)
    for evaluation in scans.get(scan_id).latest_evaluations():
        assert evaluation["check"] in text
        assert evaluation["clause"] in text
    assert report["content_sha256"] in text
    assert report["manifest"]["rulepack"]["sha256"] == pack["sha256"]
    assert report["manifest"]["inputs"][0]["sha256"]


def test_refinalizing_retains_version_one_and_uses_a_stable_content_hash(tmp_path):
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    reports = ReportService(scans, objects, pack, max_edge=1800)
    first, second = reports.finalize(scan_id), reports.finalize(scan_id)
    assert (first["version"], second["version"]) == (1, 2)
    assert first["id"] != second["id"]
    assert first["content_sha256"] == second["content_sha256"]
    assert scans.get_report(first["id"])["version"] == 1


def test_system_error_blocks_finalization(tmp_path):
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    scan = scans.get(scan_id)
    scan.evaluations[-1]["outcome"] = "SYSTEM_ERROR"
    with pytest.raises(ApiError, match="SYSTEM_ERROR") as raised:
        ReportService(scans, objects, pack, max_edge=1800).finalize(scan_id)
    assert raised.value.code == "E_CONFLICT"


def test_a_field_copy_does_not_close_the_inspection(tmp_path):
    """The single most dangerous thing a field copy could do.

    save_report used to set status=FINALIZED unconditionally, which would mean an officer
    exporting their own PDF silently locked the scan against correction, re-evaluation and
    override — destroying their own evidence by pressing Download.
    """
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    service = ReportService(scans, objects, pack, max_edge=1800)

    before = scans.get(scan_id).status
    field = service.finalize(scan_id, kind="FIELD", generated_by="officer-1")

    assert field["report_kind"] == "FIELD"
    assert scans.get(scan_id).status == before != "FINALIZED"

    finalized = service.finalize(scan_id, kind="FINALIZED", reviewed_by="reviewer-1")
    assert finalized["report_kind"] == "FINALIZED"
    assert scans.get(scan_id).status == "FINALIZED"


def test_each_kind_has_its_own_version_sequence(tmp_path):
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    service = ReportService(scans, objects, pack, max_edge=1800)
    assert service.finalize(scan_id, kind="FIELD")["version"] == 1
    assert service.finalize(scan_id, kind="FIELD")["version"] == 2
    # A field copy must not consume the finalised report's v1.
    assert service.finalize(scan_id, kind="FINALIZED")["version"] == 1


def test_a_field_copy_says_on_its_face_that_it_is_not_a_finding(tmp_path):
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    service = ReportService(scans, objects, pack, max_edge=1800)
    report = service.finalize(scan_id, kind="FIELD")
    text = "\n".join(p.text for p in
                     Document(io.BytesIO(service.download(report["id"], "docx")[0])).paragraphs)
    assert "NOT LEGALLY FINALISED" in text.upper()
    assert "Signature" not in text, "an unsigned copy must not offer a signature block"


def test_redrawing_the_report_never_changes_its_content_hash(tmp_path):
    """The hash covers the findings, not the layout.

    This is what lets the document be redesigned at all: a template change that altered a
    stored hash would silently invalidate every report already issued.
    """
    scans, objects, pack, scan_id = _evaluated(tmp_path)
    service = ReportService(scans, objects, pack, max_edge=1800)
    first = service.finalize(scan_id, kind="FIELD")
    second = service.finalize(scan_id, kind="FIELD")
    assert first["content_sha256"] == second["content_sha256"]
    assert first["manifest"]["renderer"]["template_version"] == \
        second["manifest"]["renderer"]["template_version"]


def test_evidence_is_readable_instead_of_dumped_as_json(tmp_path):
    """The old report put json.dumps(evidence) in a table cell."""
    from lmpc.server.report.sections import measured_required

    measured, required = measured_required(
        {"measured_height_mm": 1.4, "required_height_mm": 1.7})
    assert "1.4" in measured and "mm" in measured
    assert "1.7" in required and "mm" in required

    # Scoring internals are how the engine decided, not what it measured. A column headed
    # "what we measured" full of margins and token text tells a manufacturer nothing.
    assert measured_required(
        {"panel": "BACK", "score": 62, "margin": 48.9, "text": "Mfd by: X"}) == ("—", "—")
    assert measured_required({"panels_captured": ["FRONT", "BACK"]}) == ("—", "—")
    assert measured_required(None) == ("—", "—")
    assert measured_required({"nested": {"x": 1}}) == ("—", "—")

    # A limit is still a physical quantity, and belongs under "what is required".
    assert measured_required({"ratio_min": 0.6})[1] != "—"
